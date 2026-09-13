import os
import uuid
import json
from fastapi import APIRouter, Depends, UploadFile, File, Form, HTTPException
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import List, Dict
from fastapi.responses import FileResponse
import docx

from core.database import get_db
from models.template import DocumentTemplate, TemplateVersion
from services.ai_extractor import AIExtractorService

router = APIRouter()

UPLOAD_DIR = "uploads/documents"
os.makedirs(UPLOAD_DIR, exist_ok=True)

GENERATED_DIR = "uploads/generated"
os.makedirs(GENERATED_DIR, exist_ok=True)

class ExtractRequest(BaseModel):
    template_ids: List[str]

@router.post("/extract")
def extract_document_fields(
    file: UploadFile = File(...), 
    template_ids: str = Form("[]"),
    db: Session = Depends(get_db)
):
    valid_extensions = (".pdf", ".jpg", ".jpeg", ".png")
    if not any(file.filename.lower().endswith(ext) for ext in valid_extensions):
        raise HTTPException(status_code=400, detail="Only .pdf, .jpg, .jpeg, .png files are supported for extraction")
        
    doc_id = str(uuid.uuid4())
    # Strip non-ascii characters (like unicode spaces \u202f) to prevent httpx header encoding errors
    safe_filename = "".join([c for c in file.filename if c.isascii()]).replace(" ", "_")
    if not safe_filename: safe_filename = "document.pdf"
    
    filename = f"{doc_id}_{safe_filename}"
    filepath = os.path.join(UPLOAD_DIR, filename)
    
    with open(filepath, "wb") as buffer:
        content = file.file.read()
        buffer.write(content)
        
    try:
        ids_list = json.loads(template_ids)
    except:
        ids_list = []

    required_fields = set()
    try:
        if ids_list:
            versions = db.query(TemplateVersion).filter(TemplateVersion.template_id.in_(ids_list)).all()
        else:
            versions = db.query(TemplateVersion).all()
            
        for v in versions:
            if v.field_mappings:
                for k in v.field_mappings.keys():
                    required_fields.add(k)
    except Exception:
        pass
        
    # If no fields required, fallback to defaults
    if not required_fields:
        required_fields = {"deponent_name", "father_name", "address", "survey_no", "consideration_value", "registration_date"}
        
    required_fields = list(required_fields)
        
    extractor = AIExtractorService()
    extracted_values = extractor.extract_fields(filepath, required_fields)
    
    # Format the response for the frontend
    formatted_data = []
    for field, value in extracted_values.items():
        formatted_data.append({
            "label": field,
            "value": value,
            "conf": "HIGH" if value != "NOT FOUND" else "LOW",
            "confColor": "text-emerald-400 border-emerald-500/30 bg-emerald-500/10" if value != "NOT FOUND" else "text-red-400 border-red-500/30 bg-red-500/10"
        })
        
    return {
        "document_id": doc_id,
        "extracted_data": formatted_data
    }

class GenerateRequest(BaseModel):
    template_id: str
    extracted_data: Dict[str, str]

@router.post("/generate")
def generate_document(req: GenerateRequest, db: Session = Depends(get_db)):
    version = db.query(TemplateVersion).filter(TemplateVersion.template_id == req.template_id).first()
    if not version or not version.storage_key:
        raise HTTPException(status_code=404, detail="Template not found")
        
    doc = docx.Document(version.storage_key)
    
    import re
    def replace_in_paragraph(paragraph, data_dict):
        needs_replacement = False
        for key in data_dict.keys():
            if f"{{{{{key}}}}}" in paragraph.text:
                needs_replacement = True
                break
                
        if not needs_replacement:
            return

        full_text = paragraph.text
        for key, value in data_dict.items():
            placeholder = f"{{{{{key}}}}}"
            if placeholder in full_text:
                full_text = full_text.replace(placeholder, f"<<VAL_START>>{value}<<VAL_END>>")
                
        base_run = paragraph.runs[0] if paragraph.runs else None
        base_name = base_run.font.name if base_run and base_run.font else None
        base_size = base_run.font.size if base_run and base_run.font else None
        base_bold = base_run.font.bold if base_run and base_run.font else None
        base_italic = base_run.font.italic if base_run and base_run.font else None
        
        paragraph.clear()
        parts = re.split(r'(<<VAL_START>>.*?<<VAL_END>>)', full_text)
        
        for part in parts:
            if not part:
                continue
            if part.startswith("<<VAL_START>>") and part.endswith("<<VAL_END>>"):
                val = part[13:-11]
                new_run = paragraph.add_run(val)
                if base_name is not None: new_run.font.name = base_name
                if base_size is not None: new_run.font.size = base_size
                new_run.font.bold = True
            else:
                new_run = paragraph.add_run(part)
                if base_name is not None: new_run.font.name = base_name
                if base_size is not None: new_run.font.size = base_size
                if base_bold is not None: new_run.font.bold = base_bold
                if base_italic is not None: new_run.font.italic = base_italic

    # Process paragraphs
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, req.extracted_data)
                
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, req.extracted_data)
                            
    # Save the generated document
    gen_id = str(uuid.uuid4())
    filename = f"generated_{gen_id}.docx"
    filepath = os.path.join(GENERATED_DIR, filename)
    doc.save(filepath)
    
    # Read text for preview
    preview_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            preview_text.append(para.text)
            
    return {
        "document_id": gen_id,
        "download_url": f"http://127.0.0.1:8000/documents/download/{filename}",
        "preview_text": "\n\n".join(preview_text) # returning all paragraphs for preview
    }

class SaveEditedRequest(BaseModel):
    text: str
    original_filename: str
    extracted_data: Dict[str, str] = {}

@router.post("/save-edited")
def save_edited_document(req: SaveEditedRequest):
    original_filepath = os.path.join(GENERATED_DIR, req.original_filename)
    if not os.path.exists(original_filepath):
        raise HTTPException(status_code=404, detail="Original document not found")
        
    doc = docx.Document(original_filepath)
    edited_blocks = [p for p in req.text.split("\n\n") if p.strip()]
    
    import re
    block_idx = 0
    for para in doc.paragraphs:
        if para.text.strip():
            if block_idx < len(edited_blocks):
                # Retrieve base styles
                base_run = para.runs[0] if para.runs else None
                base_name = base_run.font.name if base_run and base_run.font else None
                base_size = base_run.font.size if base_run and base_run.font else None
                base_bold = base_run.font.bold if base_run and base_run.font else None
                base_italic = base_run.font.italic if base_run and base_run.font else None
                
                para.clear()
                
                # Apply bold to extracted data values
                full_text = edited_blocks[block_idx]
                for key, value in req.extracted_data.items():
                    if value and str(value) != "NOT FOUND" and len(str(value)) > 2:
                        full_text = full_text.replace(str(value), f"<<VAL_START>>{value}<<VAL_END>>")
                        
                parts = re.split(r'(<<VAL_START>>.*?<<VAL_END>>)', full_text)
                for part in parts:
                    if not part: continue
                    if part.startswith("<<VAL_START>>") and part.endswith("<<VAL_END>>"):
                        val = part[13:-11]
                        new_run = para.add_run(val)
                        if base_name is not None: new_run.font.name = base_name
                        if base_size is not None: new_run.font.size = base_size
                        new_run.font.bold = True
                    else:
                        new_run = para.add_run(part)
                        if base_name is not None: new_run.font.name = base_name
                        if base_size is not None: new_run.font.size = base_size
                        if base_bold is not None: new_run.font.bold = base_bold
                        if base_italic is not None: new_run.font.italic = base_italic
                        
                block_idx += 1
            else:
                para.text = ""
                
    while block_idx < len(edited_blocks):
        doc.add_paragraph(edited_blocks[block_idx])
        block_idx += 1
            
    gen_id = str(uuid.uuid4())
    filename = f"edited_{gen_id}.docx"
    filepath = os.path.join(GENERATED_DIR, filename)
    doc.save(filepath)
    
    return {
        "download_url": f"http://127.0.0.1:8000/documents/download/{filename}"
    }

@router.get("/download/{filename}")
def download_document(filename: str):
    filepath = os.path.join(GENERATED_DIR, filename)
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(filepath, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename=filename)
